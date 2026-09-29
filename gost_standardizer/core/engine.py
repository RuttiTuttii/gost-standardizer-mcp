from __future__ import annotations

from contextlib import contextmanager
from dataclasses import asdict
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

from gost_standardizer.core.classifier import (
    SECTION_KEYWORDS,
    classify_paragraph,
    normalize_text,
)
from gost_standardizer.core.presets import (
    BUILTIN_PROFILES,
    PRESETS,
    Preset,
    coerce_preset,
    resolve_preset,
    resolve_profile,
)
from gost_standardizer.core.rules import (
    DEFAULT_FONT_TOLERANCE_PT,
    DEFAULT_PAGE_TOLERANCE_MM,
    approx_equal,
    effective_alignment,
    effective_paragraph_metrics,
    effective_run_metrics,
    is_in_table,
    make_issue,
    paragraph_expected_metrics,
    set_page_setup,
    set_r_fonts,
    set_style_font,
)

DEFAULT_RESULT_LIMIT = 50


def resolve_input_path(raw_path: str, base_dir: Path | None = None) -> Path:
    candidate = Path(raw_path).expanduser()
    if not candidate.is_absolute():
        candidate = (base_dir or Path.cwd()) / candidate
    candidate = candidate.resolve()
    if not candidate.exists():
        raise FileNotFoundError(f"Input file not found: {candidate}")
    if candidate.suffix.lower() not in {".docx", ".docm", ".doc"}:
        raise ValueError("Only .docx, .docm, and .doc files are supported by the current implementation")
    return candidate


def make_output_path(input_path: Path, output_path: str | None = None) -> Path:
    if output_path:
        candidate = Path(output_path).expanduser()
        if not candidate.is_absolute():
            candidate = (input_path.parent / candidate).resolve()
        if candidate.suffix.lower() not in {".docx", ".docm"}:
            candidate = candidate.with_suffix(".docx")
        return candidate
    return input_path.with_name(f"{input_path.stem}_gost.docx")


def convert_legacy_doc(source: Path, temp_dir: Path) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise ValueError(
            "Legacy .doc files require LibreOffice/soffice for conversion. "
            "Install it or convert the file to .docx manually first."
        )

    temp_dir.mkdir(parents=True, exist_ok=True)
    command = [executable, "--headless", "--convert-to", "docx", "--outdir", str(temp_dir), str(source)]
    completed = subprocess.run(command, capture_output=True, text=True, check=False)
    if completed.returncode != 0:
        raise RuntimeError(
            "Failed to convert .doc to .docx with LibreOffice: "
            f"{completed.stderr.strip() or completed.stdout.strip() or 'unknown error'}"
        )

    converted = temp_dir / f"{source.stem}.docx"
    if not converted.exists():
        raise RuntimeError("LibreOffice conversion finished, but the converted .docx file was not found")
    return converted


@contextmanager
def open_document_source(raw_path: str, base_dir: Path | None = None):
    source = resolve_input_path(raw_path, base_dir=base_dir)
    if source.suffix.lower() != ".doc":
        yield source, {"source_path": str(source), "converted_from": None, "source_kind": source.suffix.lower()}
        return

    with tempfile.TemporaryDirectory(prefix="gost-standardizer-") as temp_name:
        temp_dir = Path(temp_name)
        try:
            import gost_standardizer
            conv_fn = getattr(gost_standardizer, "_convert_legacy_doc", convert_legacy_doc)
        except ImportError:
            conv_fn = convert_legacy_doc
        converted = conv_fn(source, temp_dir)
        yield converted, {

            "source_path": str(source),
            "converted_from": str(source),
            "source_kind": ".doc",
            "converted_path": str(converted),
        }


def collect_text(document: Document) -> str:
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            chunks.append(text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = normalize_text(paragraph.text)
                    if text:
                        chunks.append(text)
    return "\n".join(chunks)


def sample_paragraphs(document: Document, sample_size: int = 8) -> list[dict[str, Any]]:
    samples: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = normalize_text(paragraph.text)
        if not text:
            continue
        samples.append(
            {
                "index": index,
                "text": text,
                "style": paragraph.style.name if paragraph.style else None,
                "alignment": str(effective_alignment(paragraph)) if effective_alignment(paragraph) is not None else None,
                "kind": classify_paragraph(paragraph, index, len(samples) + 1),
            }
        )
        if len(samples) >= sample_size:
            break
    return samples


def document_statistics(document: Document) -> dict[str, Any]:
    non_empty_paragraphs = sum(1 for p in document.paragraphs if normalize_text(p.text))
    return {
        "paragraphs": len(document.paragraphs),
        "non_empty_paragraphs": non_empty_paragraphs,
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "sections": len(document.sections),
    }


def preset_scores(document: Document, source_path: Path | None = None) -> dict[str, int]:
    text = collect_text(document).lower()
    scores = {key: 0 for key in PRESETS}

    if source_path is not None:
        stem = source_path.stem.lower()
        if any(token in stem for token in ("tz", "тз", "technical", "spec")):
            scores["technical"] += 2
        if any(token in stem for token in ("report", "отчет", "otchet", "пз", "poyasnit")):
            scores["report"] += 2
        if any(token in stem for token in ("order", "приказ", "reglament", "instruction")):
            scores["office"] += 2

    for key, keywords in SECTION_KEYWORDS.items():
        for keyword in keywords:
            if keyword in text:
                scores[key] += 1
    return scores


def detect_preset(path: str, preset_name: str | None = None) -> dict[str, Any]:
    with open_document_source(path) as (source, source_meta):
        document = Document(str(source))
        scores = preset_scores(document, source)
        guessed = max(scores, key=lambda k: scores[k]) if max(scores.values(), default=0) > 0 else "report"
        selected = preset_name or guessed
        preset = resolve_preset(selected)
        return {
            "path": str(source_meta["source_path"]),
            "source_kind": source_meta["source_kind"],
            "converted_from": source_meta.get("converted_from"),
            "requested_preset": preset_name,
            "preset": asdict(preset),
            "guessed_preset": guessed,
            "scores": scores,
        }


def explain_preset(
    path_or_preset: str | None = None,
    preset_name: str | None = None,
    name: str | None = None,
) -> dict[str, Any]:
    target_preset = preset_name or name
    if path_or_preset and (Path(path_or_preset).exists() or path_or_preset.endswith((".docx", ".docm", ".doc"))):
        with open_document_source(path_or_preset) as (source, source_meta):
            document = Document(str(source))
            scores = preset_scores(document, source)
            guessed = max(scores, key=lambda k: scores[k]) if max(scores.values(), default=0) > 0 else "report"
            selected = target_preset or guessed
            preset = resolve_preset(selected)

            stem = source.stem.lower()
            signals: list[dict[str, Any]] = []
            if any(token in stem for token in ("tz", "тз", "technical", "spec")):
                signals.append({"signal": "filename", "token": stem, "preset": "technical", "weight": 2})
            if any(token in stem for token in ("report", "отчет", "otchet", "пз", "poyasnit")):
                signals.append({"signal": "filename", "token": stem, "preset": "report", "weight": 2})
            if any(token in stem for token in ("order", "приказ", "reglament", "instruction")):
                signals.append({"signal": "filename", "token": stem, "preset": "office", "weight": 2})

            lowered = collect_text(document).lower()
            for key, keywords in SECTION_KEYWORDS.items():
                for keyword in keywords:
                    if keyword in lowered:
                        signals.append({"signal": "keyword", "token": keyword, "preset": key, "weight": 1})

            return {
                "path": str(source_meta["source_path"]),
                "source_kind": source_meta["source_kind"],
                "converted_from": source_meta.get("converted_from"),
                "requested_preset": target_preset,
                "preset": asdict(preset),
                "guessed_preset": guessed,
                "scores": scores,
                "signals": signals,
            }

    preset = resolve_preset(path_or_preset or target_preset)
    return {
        "key": preset.key,
        "title": preset.title,
        "description": preset.description,
        "preset": asdict(preset),
    }



def validate_page_setup(document: Document, preset: Preset) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    matches: list[dict[str, Any]] = []

    for section_index, section in enumerate(document.sections):
        actual = {
            "page_width_mm": round(section.page_width.mm, 2),
            "page_height_mm": round(section.page_height.mm, 2),
            "left_margin_mm": round(section.left_margin.mm, 2),
            "right_margin_mm": round(section.right_margin.mm, 2),
            "top_margin_mm": round(section.top_margin.mm, 2),
            "bottom_margin_mm": round(section.bottom_margin.mm, 2),
        }
        expected = {
            "page_width_mm": preset.page_width_mm,
            "page_height_mm": preset.page_height_mm,
            "left_margin_mm": preset.margin_left_mm,
            "right_margin_mm": preset.margin_right_mm,
            "top_margin_mm": preset.margin_top_mm,
            "bottom_margin_mm": preset.margin_bottom_mm,
        }

        for key in ("page_width_mm", "page_height_mm"):
            if not approx_equal(actual[key], expected[key], DEFAULT_PAGE_TOLERANCE_MM):
                issues.append(
                    make_issue(
                        "page.size.mismatch",
                        f"Section {section_index + 1} {key} is {actual[key]} mm, expected {expected[key]} mm",
                        evidence={"section": section_index, "actual": actual, "expected": expected},
                    )
                )
                break
        else:
            matches.append(
                {
                    "rule_id": "page.size.mismatch",
                    "section": section_index,
                    "status": "match",
                    "actual": actual,
                    "expected": expected,
                }
            )

        for key, rule_id in (
            ("left_margin_mm", "page.margin.left"),
            ("right_margin_mm", "page.margin.right"),
            ("top_margin_mm", "page.margin.top"),
            ("bottom_margin_mm", "page.margin.bottom"),
        ):
            if not approx_equal(actual[key], expected[key], DEFAULT_PAGE_TOLERANCE_MM):
                issues.append(
                    make_issue(
                        rule_id,
                        f"Section {section_index + 1} {key} is {actual[key]} mm, expected {expected[key]} mm",
                        evidence={"section": section_index, "actual": actual, "expected": expected},
                    )
                )
            else:
                matches.append(
                    {
                        "rule_id": rule_id,
                        "section": section_index,
                        "status": "match",
                        "actual": actual[key],
                        "expected": expected[key],
                    }
                )

    return issues, matches


def validate_paragraphs(
    document: Document,
    preset: Preset,
    *,
    aggressive: bool = False,
    limit: int = DEFAULT_RESULT_LIMIT,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    matches: list[dict[str, Any]] = []
    non_empty_seen = 0

    for index, paragraph in enumerate(document.paragraphs):
        text = normalize_text(paragraph.text)
        if not text:
            continue

        non_empty_seen += 1
        kind = classify_paragraph(paragraph, index, non_empty_seen)
        expected = paragraph_expected_metrics(kind, preset)
        actual = effective_paragraph_metrics(paragraph)

        if kind in {"title", "heading", "caption", "list", "body"}:
            matches.append(
                {
                    "index": index,
                    "kind": kind,
                    "text": text[:160],
                    "status": "analyzed",
                }
            )

        if is_in_table(paragraph):
            continue

        # Check alignment
        if actual["alignment"] != expected["alignment"]:
            issues.append(
                make_issue(
                    "paragraph.alignment.mismatch",
                    f"Paragraph {index + 1} ({kind}) alignment does not match profile.",
                    evidence={"index": index, "kind": kind, "actual": str(actual["alignment"])},
                )
            )

        # Check first line indent
        if not approx_equal(actual["first_line_indent_mm"], expected["first_line_indent_mm"], 1.5):
            issues.append(
                make_issue(
                    "paragraph.indent.mismatch",
                    f"Paragraph {index + 1} ({kind}) indent is {actual['first_line_indent_mm']} mm, expected {expected['first_line_indent_mm']} mm.",
                    evidence={"index": index, "kind": kind},
                )
            )

        # Check runs fonts
        for run in paragraph.runs:
            rm = effective_run_metrics(run, paragraph)
            if rm["name"] and rm["name"].lower() != preset.body_font_name.lower():
                issues.append(
                    make_issue(
                        "run.font.family.mismatch",
                        f"Paragraph {index + 1} run uses font '{rm['name']}', expected '{preset.body_font_name}'.",
                        evidence={"index": index, "actual_font": rm["name"]},
                    )
                )
                break

    return issues, matches


def analyze_document(
    path: str,
    preset_name: str | None = None,
    profile_name: str | None = None,
    sample_size: int = 8,
    aggressive: bool = False,
    limit: int = DEFAULT_RESULT_LIMIT,
) -> dict[str, Any]:
    with open_document_source(path) as (source, source_meta):
        document = Document(str(source))
        target_profile = profile_name or preset_name
        profile = resolve_profile(target_profile)
        preset = coerce_preset(profile["preset"])

        issues: list[dict[str, Any]] = []
        page_issues, page_matches = validate_page_setup(document, preset)
        paragraph_issues, paragraph_matches = validate_paragraphs(document, preset, aggressive=aggressive, limit=limit)
        issues.extend(page_issues)
        issues.extend(paragraph_issues)

        recommendations = list(
            dict.fromkeys(issue["recommendation"] for issue in issues if issue.get("recommendation"))
        )

        severity_counts: dict[str, int] = {}
        for issue in issues:
            severity_counts[issue["severity"]] = severity_counts.get(issue["severity"], 0) + 1

        return {
            "kind": "analysis",
            "path": str(source_meta["source_path"]),
            "source_kind": source_meta["source_kind"],
            "converted_from": source_meta.get("converted_from"),
            "profile": profile,
            "preset": asdict(preset),
            "statistics": document_statistics(document),
            "summary": {
                "issues": len(issues),
                "errors": severity_counts.get("error", 0),
                "warnings": severity_counts.get("warning", 0),
                "info": severity_counts.get("info", 0),
                "auto_fixable": sum(1 for issue in issues if issue.get("auto_fixable")),
            },
            "page_matches": page_matches,
            "paragraph_matches": paragraph_matches,
            "issues": issues[:limit],
            "recommendations": recommendations,
            "sample_paragraphs": sample_paragraphs(document, sample_size),
        }


def inspect_document(path: str, sample_size: int = 8) -> dict[str, Any]:
    analysis = analyze_document(path, sample_size=sample_size)
    return {
        "path": analysis["path"],
        "source_kind": analysis["source_kind"],
        "converted_from": analysis["converted_from"],
        "preset_guess": analysis["profile"]["key"],
        "profile": analysis["profile"],
        "preset": analysis["preset"],
        "statistics": analysis["statistics"],
        "deviations": [issue["message"] for issue in analysis["issues"]],
        "issues": analysis["issues"],
        "recommendations": analysis["recommendations"],
        "sample_paragraphs": analysis["sample_paragraphs"],
    }


def validate_document(
    path: str,
    preset_name: str | None = None,
    profile_name: str | None = None,
    aggressive: bool = False,
    sample_size: int = 8,
) -> dict[str, Any]:
    target_profile = profile_name or preset_name
    analysis = analyze_document(path, profile_name=target_profile, sample_size=sample_size, aggressive=aggressive)
    return {
        "kind": "validation",
        "path": analysis["path"],
        "source_kind": analysis["source_kind"],
        "converted_from": analysis["converted_from"],
        "profile": analysis["profile"],
        "preset": analysis["preset"],
        "statistics": analysis["statistics"],
        "summary": analysis["summary"],
        "issues": analysis["issues"],
        "recommendations": analysis["recommendations"],
        "sample_paragraphs": analysis["sample_paragraphs"],
    }


def compare_to_preset(
    path: str,
    preset_name: str | None = None,
    profile_name: str | None = None,
    aggressive: bool = False,
    sample_size: int = 8,
) -> dict[str, Any]:
    target = profile_name or preset_name or "report"
    profile = resolve_profile(target)
    analysis = analyze_document(path, profile_name=target, sample_size=sample_size, aggressive=aggressive)
    issues = analysis["issues"]
    matches = analysis["page_matches"] + analysis["paragraph_matches"]
    return {
        "kind": "comparison",
        "path": analysis["path"],
        "source_kind": analysis["source_kind"],
        "converted_from": analysis["converted_from"],
        "profile": profile,
        "preset": analysis["preset"],
        "summary": {
            "matches": len(matches),
            "differences": len(issues),
            "auto_fixable": sum(1 for issue in issues if issue.get("auto_fixable")),
            "errors": analysis["summary"]["errors"],
            "warnings": analysis["summary"]["warnings"],
        },
        "differences": issues,
        "matches": matches,
        "statistics": analysis["statistics"],
        "recommendations": analysis["recommendations"],
        "sample_paragraphs": analysis["sample_paragraphs"],
    }


def apply_paragraph_format(paragraph: Any, kind: str, preset: Preset, inside_table: bool = False) -> None:
    fmt = paragraph.paragraph_format
    if kind == "empty":
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        return

    if inside_table:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.first_line_indent = Mm(0)
        fmt.left_indent = Mm(0)
        fmt.right_indent = Mm(0)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        return

    if kind == "title":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.first_line_indent = Mm(0)
        fmt.left_indent = Mm(0)
        fmt.right_indent = Mm(0)
        fmt.space_before = Pt(preset.title_space_before_pt)
        fmt.space_after = Pt(preset.title_space_after_pt)
        fmt.line_spacing = 1.0
        return

    if kind == "heading":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.first_line_indent = Mm(0)
        fmt.left_indent = Mm(0)
        fmt.right_indent = Mm(0)
        fmt.space_before = Pt(preset.heading_space_before_pt)
        fmt.space_after = Pt(preset.heading_space_after_pt)
        fmt.line_spacing = 1.0
        return

    if kind == "caption":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.first_line_indent = Mm(0)
        fmt.left_indent = Mm(0)
        fmt.right_indent = Mm(0)
        fmt.space_before = Pt(preset.caption_space_before_pt)
        fmt.space_after = Pt(preset.caption_space_after_pt)
        fmt.line_spacing = 1.0
        return

    if kind == "list":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Mm(8)
        fmt.first_line_indent = Mm(0)
        fmt.right_indent = Mm(0)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = preset.body_line_spacing
        return

    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Mm(preset.body_first_line_indent_mm)
    fmt.left_indent = Mm(0)
    fmt.right_indent = Mm(0)
    fmt.space_before = Pt(preset.body_space_before_pt)
    fmt.space_after = Pt(preset.body_space_after_pt)
    fmt.line_spacing = preset.body_line_spacing


def apply_run_format(run: Any, kind: str, preset: Preset, *, preserve_inline_styles: bool = False) -> None:
    if not run.text:
        return

    if kind == "title":
        run.font.name = preset.body_font_name
        run.font.size = Pt(preset.title_font_size_pt)
        run.font.bold = True
        run.font.italic = False
        set_r_fonts(run, preset.body_font_name)
        return

    if kind == "heading":
        run.font.name = preset.body_font_name
        run.font.size = Pt(preset.heading_font_size_pt)
        run.font.bold = True
        run.font.italic = False
        set_r_fonts(run, preset.body_font_name)
        return

    if kind == "caption":
        run.font.name = preset.body_font_name
        run.font.size = Pt(preset.table_font_size_pt)
        run.font.bold = False
        run.font.italic = True
        set_r_fonts(run, preset.body_font_name)
        return

    size = preset.table_font_size_pt if kind == "list" else preset.body_font_size_pt
    run.font.name = preset.body_font_name
    run.font.size = Pt(size)
    if not preserve_inline_styles:
        run.font.bold = False
        run.font.italic = False
    set_r_fonts(run, preset.body_font_name)


def apply_style_defaults(document: Document, preset: Preset) -> None:
    for style_name in [
        "Normal",
        "Body Text",
        "List Paragraph",
        "Caption",
        "Title",
        "Subtitle",
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ]:
        try:
            style = document.styles[style_name]
        except KeyError:
            continue

        if style_name in {"Heading 1", "Heading 2", "Heading 3"}:
            set_style_font(style, preset.body_font_name, preset.heading_font_size_pt, True)
            continue
        if style_name == "Title":
            set_style_font(style, preset.body_font_name, preset.title_font_size_pt, True)
            continue
        if style_name == "Caption":
            set_style_font(style, preset.body_font_name, preset.table_font_size_pt, False)
            continue
        set_style_font(style, preset.body_font_name, preset.body_font_size_pt, False)

        if style.paragraph_format is not None:
            fmt = style.paragraph_format
            fmt.space_before = Pt(preset.body_space_before_pt)
            fmt.space_after = Pt(preset.body_space_after_pt)
            fmt.line_spacing = preset.body_line_spacing
            if style_name == "List Paragraph":
                fmt.first_line_indent = Mm(0)
                fmt.left_indent = Mm(8)
            else:
                fmt.first_line_indent = Mm(preset.body_first_line_indent_mm)
                fmt.left_indent = Mm(0)


def standardize_document(
    path: str,
    output_path: str | None = None,
    preset_name: str | None = None,
    profile_name: str | None = None,
    overwrite: bool = False,
    aggressive: bool = False,
    fix_page_setup: bool = True,
    fix_styles: bool = True,
    fix_paragraphs: bool = True,
    fix_tables: bool = True,
) -> dict[str, Any]:
    target = profile_name or preset_name
    if target:
        profile = resolve_profile(target)
        preset = coerce_preset(profile["preset"])
    else:
        profile = dict(BUILTIN_PROFILES["report"])
        preset = coerce_preset(profile["preset"])

    with open_document_source(path) as (source, source_meta):
        document = Document(str(source))
        if fix_page_setup:
            set_page_setup(document, preset)
        if fix_styles:
            apply_style_defaults(document, preset)

        changed_paragraphs = 0
        non_empty_seen = 0
        if fix_paragraphs:
            for index, paragraph in enumerate(document.paragraphs):
                text = normalize_text(paragraph.text)
                if not text:
                    continue
                non_empty_seen += 1
                kind = classify_paragraph(paragraph, index, non_empty_seen)
                apply_paragraph_format(paragraph, kind, preset)
                preserve_inline = not aggressive and kind == "body"
                for run in paragraph.runs:
                    apply_run_format(run, kind, preset, preserve_inline_styles=preserve_inline)
                changed_paragraphs += 1

        changed_tables = 0
        if fix_tables:
            for table in document.tables:
                changed_tables += 1
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            apply_paragraph_format(paragraph, "body", preset, inside_table=True)
                            for run in paragraph.runs:
                                run.font.name = preset.body_font_name
                                run.font.size = Pt(preset.table_font_size_pt)
                                set_r_fonts(run, preset.body_font_name)

        original_input_path = Path(source_meta["source_path"])
        dest = make_output_path(original_input_path, output_path)
        if dest.exists() and not overwrite:
            raise FileExistsError(f"Destination file already exists: {dest}")

        dest.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(dest))

        return {
            "kind": "standardization",
            "path": str(original_input_path),
            "output_path": str(dest),
            "source_kind": source_meta["source_kind"],
            "converted_from": source_meta.get("converted_from"),
            "profile": profile,
            "preset": asdict(preset),
            "statistics": document_statistics(document),
            "changes": {
                "page_setup": fix_page_setup,
                "styles": fix_styles,
                "paragraphs": changed_paragraphs,
                "tables": changed_tables,
            },
        }
