from __future__ import annotations

import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import MagicMock, patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from gost_standardizer.converter import html_markdown
from gost_standardizer.converter.html_markdown import (
    _SimpleHtmlToMarkdownParser,
    convert_html_to_markdown,
)
from gost_standardizer.core import classifier, engine, presets, rules
from gost_standardizer.core.classifier import classify_paragraph, has_numbering, normalize_text
from gost_standardizer.core.engine import (
    analyze_document,
    apply_paragraph_format,
    apply_run_format,
    apply_style_defaults,
    collect_text,
    convert_legacy_doc,
    detect_preset,
    explain_preset,
    make_output_path,
    preset_scores,
    resolve_input_path,
    sample_paragraphs,
    standardize_document,
    validate_page_setup,
    validate_paragraphs,
)
from gost_standardizer.core.presets import (
    Preset,
    coerce_preset,
    list_presets,
    list_profiles,
    load_profile,
    profile_from_payload,
    resolve_preset,
)
from gost_standardizer.core.rules import (
    effective_alignment,
    effective_paragraph_metrics,
    is_in_table,
    paragraph_expected_metrics,
    set_style_font,
)
from gost_standardizer.models.document import DocumentStatus, NormDocument
from gost_standardizer.models.profiles import DocumentProfile, FontProfile, MarginsProfile, SpacingProfile
from gost_standardizer.models.validation import InspectionReport, StandardizationResult, ValidationIssue
from gost_standardizer.server.transport import StdioTransport


class EngineExtendedTests(unittest.TestCase):
    def test_resolve_input_path_cases(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            doc_file = base / "valid.docx"
            doc_file.write_bytes(b"")
            txt_file = base / "invalid.txt"
            txt_file.write_bytes(b"")

            # Relative path resolution with base_dir
            resolved = resolve_input_path("valid.docx", base_dir=base)
            self.assertEqual(resolved, doc_file.resolve())

            # File not found
            with self.assertRaises(FileNotFoundError):
                resolve_input_path("nonexistent.docx", base_dir=base)

            # Invalid extension
            with self.assertRaises(ValueError) as ctx:
                resolve_input_path(str(txt_file))
            self.assertIn("Only .docx, .docm, and .doc files", str(ctx.exception))

    def test_make_output_path_cases(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            in_file = base / "sub" / "input.docx"

            # Relative output path
            out1 = make_output_path(in_file, "custom_out.docx")
            self.assertEqual(out1, (in_file.parent / "custom_out.docx").resolve())

            # Output path without docx extension
            out2 = make_output_path(in_file, "custom_out.pdf")
            self.assertEqual(out2.suffix, ".docx")

    def test_convert_legacy_doc_scenarios(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            src_doc = base / "legacy.doc"
            src_doc.write_bytes(b"dummy")
            out_dir = base / "out"

            # 1. Non-zero return code
            mock_proc = MagicMock()
            mock_proc.returncode = 1
            mock_proc.stderr = "Conversion crash"
            mock_proc.stdout = ""
            with patch("shutil.which", return_value="/usr/bin/soffice"):
                with patch("subprocess.run", return_value=mock_proc):
                    with self.assertRaises(RuntimeError) as ctx:
                        convert_legacy_doc(src_doc, out_dir)
                    self.assertIn("Failed to convert .doc to .docx", str(ctx.exception))

            # 2. Zero return code, but converted file not created
            mock_proc.returncode = 0
            with patch("shutil.which", return_value="/usr/bin/soffice"):
                with patch("subprocess.run", return_value=mock_proc):
                    with self.assertRaises(RuntimeError) as ctx:
                        convert_legacy_doc(src_doc, out_dir)
                    self.assertIn("converted .docx file was not found", str(ctx.exception))

            # 3. Successful conversion
            def create_docx(*args, **kwargs):
                (out_dir / "legacy.docx").write_bytes(b"docx content")
                return mock_proc

            with patch("shutil.which", return_value="/usr/bin/soffice"):
                with patch("subprocess.run", side_effect=create_docx):
                    converted = convert_legacy_doc(src_doc, out_dir)
                    self.assertTrue(converted.exists())

    def test_collect_text_and_sample_paragraphs_with_tables_and_limits(self) -> None:
        doc = Document()
        doc.add_paragraph("")  # empty paragraph
        doc.add_paragraph("Первый абзац документа")
        doc.add_paragraph("Второй абзац документа")

        table = doc.add_table(rows=1, cols=1)
        cell_p = table.rows[0].cells[0].paragraphs[0]
        cell_p.text = "Текст внутри ячейки таблицы"

        full_text = collect_text(doc)
        self.assertIn("Первый абзац документа", full_text)
        self.assertIn("Текст внутри ячейки таблицы", full_text)

        # sample_paragraphs with sample_size = 1 (triggers limit break)
        samples = sample_paragraphs(doc, sample_size=1)
        self.assertEqual(len(samples), 1)
        self.assertEqual(samples[0]["text"], "Первый абзац документа")

    def test_preset_scores_and_explain_preset_signals(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)

            # Test technical filename stem token
            tz_file = base / "tz_project.docx"
            doc_tz = Document()
            doc_tz.add_paragraph("Техническое задание на разработку")
            doc_tz.save(tz_file)

            scores_tz = preset_scores(doc_tz, tz_file)
            self.assertGreater(scores_tz["technical"], 0)

            # Test office filename stem token
            order_file = base / "order_123.docx"
            doc_order = Document()
            doc_order.add_paragraph("Приказ по основной деятельности")
            doc_order.save(order_file)

            scores_order = preset_scores(doc_order, order_file)
            self.assertGreater(scores_order["office"], 0)

            # detect_preset
            detected = detect_preset(str(tz_file))
            self.assertEqual(detected["guessed_preset"], "technical")

            # explain_preset with document path having tz and order signals
            exp_tz = explain_preset(str(tz_file))
            self.assertEqual(exp_tz["guessed_preset"], "technical")
            signal_presets = [s["preset"] for s in exp_tz["signals"]]
            self.assertIn("technical", signal_presets)

            exp_order = explain_preset(str(order_file))
            self.assertEqual(exp_order["guessed_preset"], "office")

            # explain_preset with preset name only
            exp_name = explain_preset("office")
            self.assertEqual(exp_name["key"], "office")

    def test_validate_page_setup_matching_and_paragraphs(self) -> None:
        doc = Document()
        sec = doc.sections[0]
        preset = resolve_preset("report")
        sec.page_width = Mm(preset.page_width_mm)
        sec.page_height = Mm(preset.page_height_mm)
        sec.left_margin = Mm(preset.margin_left_mm)
        sec.right_margin = Mm(preset.margin_right_mm)
        sec.top_margin = Mm(preset.margin_top_mm)
        sec.bottom_margin = Mm(preset.margin_bottom_mm)

        issues, matches = validate_page_setup(doc, preset)
        self.assertEqual(len(issues), 0)
        self.assertGreater(len(matches), 0)

    def test_validate_paragraphs_table_skipping_and_empty(self) -> None:
        doc = Document()
        preset = resolve_preset("report")
        doc.add_paragraph("")  # empty paragraph

        p_body = doc.add_paragraph("Обычный абзац текста.")
        p_body.alignment = WD_ALIGN_PARAGRAPH.LEFT  # Mismatch with JUSTIFY

        table = doc.add_table(rows=1, cols=1)
        p_in_tbl = table.rows[0].cells[0].paragraphs[0]
        p_in_tbl.text = "Текст в таблице"

        issues, matches = validate_paragraphs(doc, preset)
        # Table paragraph should be skipped from issues
        self.assertTrue(any(i["rule_id"] == "paragraph.alignment.mismatch" for i in issues))

    def test_apply_paragraph_format_and_run_format_all_kinds(self) -> None:
        doc = Document()
        preset = resolve_preset("report")

        # 1. empty kind
        p_empty = doc.add_paragraph("")
        apply_paragraph_format(p_empty, "empty", preset)
        self.assertEqual(p_empty.paragraph_format.space_before, Pt(0))

        # 2. inside_table
        p_tbl = doc.add_paragraph("В таблице")
        apply_paragraph_format(p_tbl, "body", preset, inside_table=True)
        self.assertEqual(p_tbl.alignment, WD_ALIGN_PARAGRAPH.LEFT)

        # 3. caption
        p_cap = doc.add_paragraph("Рисунок 1 - Архитектура")
        apply_paragraph_format(p_cap, "caption", preset)
        self.assertEqual(p_cap.alignment, WD_ALIGN_PARAGRAPH.CENTER)

        # 4. list
        p_list = doc.add_paragraph("- пункт списка")
        apply_paragraph_format(p_list, "list", preset)
        self.assertAlmostEqual(p_list.paragraph_format.left_indent.mm, 8.0, places=1)

        # 5. run format: empty run text
        r_empty = p_empty.add_run("")
        apply_run_format(r_empty, "title", preset)

        # 6. run format: caption
        r_cap = p_cap.runs[0]
        apply_run_format(r_cap, "caption", preset)
        self.assertTrue(r_cap.font.italic)
        self.assertFalse(r_cap.font.bold)

        # 7. run format: body with preserve_inline_styles=False
        p_body = doc.add_paragraph()
        r_body = p_body.add_run("Текст")
        r_body.font.bold = True
        apply_run_format(r_body, "body", preset, preserve_inline_styles=False)
        self.assertFalse(r_body.font.bold)

    def test_apply_style_defaults_missing_style(self) -> None:
        preset = resolve_preset("report")
        mock_doc = MagicMock()
        mock_doc.styles.__getitem__.side_effect = KeyError("Style not found")
        # Should catch KeyError and continue cleanly without raising
        apply_style_defaults(mock_doc, preset)

    def test_standardize_document_tables_and_exists_error(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            in_file = base / "doc_with_table.docx"
            out_file = base / "out_exists.docx"

            doc = Document()
            doc.add_paragraph("")  # empty paragraph
            table = doc.add_table(rows=1, cols=1)
            table.rows[0].cells[0].paragraphs[0].text = "Ячейка таблицы"
            doc.save(in_file)

            # Pre-create output file to test overwrite=False
            out_file.write_bytes(b"existing")

            with self.assertRaises(FileExistsError):
                standardize_document(str(in_file), output_path=str(out_file), overwrite=False)

            # Standardize with overwrite=True
            res = standardize_document(
                str(in_file),
                output_path=str(out_file),
                overwrite=True,
                fix_tables=True,
            )
            self.assertEqual(res["changes"]["tables"], 1)

    def test_classifier_all_branches(self) -> None:
        doc = Document()

        # 1. empty text
        p_empty = doc.add_paragraph("   ")
        self.assertEqual(classify_paragraph(p_empty, 0, 1), "empty")

        # 2. style name branches using mock paragraph
        p_mock = MagicMock()
        p_mock.text = "Некоторый текст"

        p_mock.style.name = "Document Title"
        self.assertEqual(classify_paragraph(p_mock, 0, 1), "title")

        p_mock.style.name = "Heading 1"
        self.assertEqual(classify_paragraph(p_mock, 0, 1), "heading")

        p_mock.style.name = "Table Caption"
        self.assertEqual(classify_paragraph(p_mock, 0, 1), "caption")

        # 3. numbering property
        p_num = doc.add_paragraph("Пункт с нумерацией")
        p_num.style = None
        p_pr = p_num._p.get_or_add_pPr()
        p_pr.append(OxmlElement("w:numPr"))
        self.assertTrue(has_numbering(p_num))
        self.assertEqual(classify_paragraph(p_num, 5, 5), "list")

        # 4. non_empty_index <= 3 and CENTER alignment -> title
        p_title_center = doc.add_paragraph("Название документа без точки")
        p_title_center.style = None
        p_title_center.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.assertEqual(classify_paragraph(p_title_center, 2, 2), "title")

        # 5. Uppercase heading
        p_upper = doc.add_paragraph("ОБЩИЕ ПОЛОЖЕНИЯ")
        p_upper.style = None
        p_upper.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.assertEqual(classify_paragraph(p_upper, 5, 5), "heading")

        # 6. Caption startswith рисунок / таблица
        p_fig = doc.add_paragraph("Рисунок 2. Схема взаимодействия")
        p_fig.style = None
        self.assertEqual(classify_paragraph(p_fig, 6, 6), "caption")

        p_tab = doc.add_paragraph("Таблица 1. Показатели")
        p_tab.style = None
        self.assertEqual(classify_paragraph(p_tab, 7, 7), "caption")

        # 7. Short capitalized sentence without period -> heading
        p_short = doc.add_paragraph("Краткий заголовок раздела")
        p_short.style = None
        self.assertEqual(classify_paragraph(p_short, 8, 8), "heading")

    def test_presets_and_profiles_edge_cases(self) -> None:
        # 1. coerce_preset missing fields -> ValueError
        with self.assertRaises(ValueError) as ctx:
            coerce_preset({"key": "incomplete"})
        self.assertIn("Profile preset is missing fields", str(ctx.exception))

        # 2. profile_from_payload with Preset object
        p_obj = resolve_preset("report")
        payload_obj = profile_from_payload({"preset": p_obj, "key": "test_obj"})
        self.assertEqual(payload_obj["key"], "test_obj")

        # 3. profile_from_payload without preset mapping -> ValueError
        with self.assertRaises(ValueError):
            profile_from_payload({"preset": "not_a_dict"})

        # 4. resolve_preset unknown -> ValueError
        with self.assertRaises(ValueError):
            resolve_preset("nonexistent_preset")

        # 5. load_profile unknown -> ValueError
        with self.assertRaises(ValueError):
            load_profile("nonexistent_profile")

        # 6. load_profile from file path
        with tempfile.TemporaryDirectory() as td:
            p_file = Path(td) / "custom.json"
            p_file.write_text(json.dumps(payload_obj), encoding="utf-8")
            loaded = load_profile(str(p_file))
            self.assertEqual(loaded["source"], "file")

            # 7. list_profiles with corrupted file in PROFILES_DIR
            corrupt_file = Path(td) / "corrupt.json"
            corrupt_file.write_text("NOT JSON", encoding="utf-8")
            with patch("gost_standardizer.core.presets.PROFILES_DIR", Path(td)):
                profs = list_profiles()
                # Contains built-in plus the valid custom.json, corrupt was skipped
                keys = [p["key"] for p in profs]
                self.assertIn("test_obj", keys)

        # 8. explain_preset in presets module
        exp = presets.explain_preset("office")
        self.assertEqual(exp["key"], "office")

    def test_rules_helpers_and_metrics(self) -> None:
        doc = Document()
        p = doc.add_paragraph("Тест правил")

        # 1. set_style_font when r_fonts is None
        style = doc.styles.add_style("CustomStyle", 1)
        r_pr = style._element.get_or_add_rPr()
        # Ensure rFonts is not present
        if r_pr.find(qn("w:rFonts")) is not None:
            r_pr.remove(r_pr.find(qn("w:rFonts")))
        set_style_font(style, "Arial", 12, True)
        self.assertIsNotNone(r_pr.find(qn("w:rFonts")))

        # 2. is_in_table on table paragraph
        tbl = doc.add_table(rows=1, cols=1)
        p_cell = tbl.rows[0].cells[0].paragraphs[0]
        self.assertTrue(is_in_table(p_cell))
        self.assertFalse(is_in_table(p))

        # 3. effective_alignment from style
        p.alignment = None
        p.style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.assertEqual(effective_alignment(p), WD_ALIGN_PARAGRAPH.RIGHT)

        # 4. effective_paragraph_metrics style fallback
        metrics = effective_paragraph_metrics(p)
        self.assertEqual(metrics["alignment"], WD_ALIGN_PARAGRAPH.RIGHT)

        # 5. paragraph_expected_metrics for title, caption, list
        preset = resolve_preset("report")
        m_title = paragraph_expected_metrics("title", preset)
        self.assertEqual(m_title["alignment"], WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(m_title["bold"])

        m_cap = paragraph_expected_metrics("caption", preset)
        self.assertEqual(m_cap["alignment"], WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(m_cap["italic"])

        m_list = paragraph_expected_metrics("list", preset)
        self.assertEqual(m_list["alignment"], WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(m_list["left_indent_mm"], 8.0)

    def test_models_dataclasses(self) -> None:
        # DocumentStatus.from_str branches
        self.assertEqual(DocumentStatus.from_str(None), DocumentStatus.UNKNOWN)
        self.assertEqual(DocumentStatus.from_str(""), DocumentStatus.UNKNOWN)
        self.assertEqual(DocumentStatus.from_str("не вступил в силу"), DocumentStatus.NOT_YET_ACTIVE)
        self.assertEqual(DocumentStatus.from_str("утратил силу"), DocumentStatus.CANCELLED)
        self.assertEqual(DocumentStatus.from_str("какой-то другой статус"), DocumentStatus.UNKNOWN)

        # NormDocument empty clean_number auto-population and to_dict
        nd = NormDocument(id="1", designation="ГОСТ 1234-2020", clean_number="", title="Тест")
        self.assertEqual(nd.clean_number, "1234-2020")
        nd_dict = nd.to_dict()
        self.assertEqual(nd_dict["clean_number"], "1234-2020")

        # DocumentProfile to_dict and from_dict
        prof = DocumentProfile(name="custom_prof")
        prof_dict = prof.to_dict()
        prof_restored = DocumentProfile.from_dict(prof_dict)
        self.assertEqual(prof_restored.name, "custom_prof")

        # Validation models to_dict
        vi = ValidationIssue(kind="test", message="msg")
        vi_dict = vi.to_dict()
        self.assertEqual(vi_dict["kind"], "test")

        ir = InspectionReport(
            path="test.docx",
            format="docx",
            paragraphs_count=10,
            tables_count=1,
            suggested_preset="report",
            issues=[vi],
        )
        ir_dict = ir.to_dict()
        self.assertEqual(ir_dict["issues"][0]["kind"], "test")

        sr = StandardizationResult(
            source_path="src.docx",
            output_path="out.docx",
            preset="report",
            changed_paragraphs=5,
            changed_tables=1,
            margins_adjusted=True,
        )
        sr_dict = sr.to_dict()
        self.assertEqual(sr_dict["preset"], "report")

    def test_converter_html_to_markdown_extended(self) -> None:
        # Empty HTML
        self.assertEqual(convert_html_to_markdown(""), "")

        # _native_h2m with result.content
        mock_native = MagicMock()
        mock_res = MagicMock()
        mock_res.content = "Converted markdown content"
        mock_native.convert.return_value = mock_res
        with patch.object(html_markdown, "_native_h2m", mock_native):
            self.assertEqual(convert_html_to_markdown("<p>test</p>"), "Converted markdown content")

        # _native_h2m raising exception, falls through to Strategy 2 (CLI)
        mock_native.convert.side_effect = RuntimeError("native failed")
        mock_proc = MagicMock()
        mock_proc.returncode = 0
        mock_proc.stdout = b"# Markdown from CLI\n"
        with patch.object(html_markdown, "_native_h2m", mock_native):
            with patch("shutil.which", return_value="/usr/bin/html-to-markdown"):
                with patch("os.path.isfile", return_value=True):
                    with patch("os.access", return_value=True):
                        with patch("subprocess.run", return_value=mock_proc):
                            res = convert_html_to_markdown("<h1>Test</h1>")
                            self.assertEqual(res, "# Markdown from CLI")

        # CLI returning non-zero returncode falls through to Strategy 3 (Pure python)
        mock_proc.returncode = 1
        with patch.object(html_markdown, "_native_h2m", None):
            with patch("shutil.which", return_value="/usr/bin/html-to-markdown"):
                with patch("os.path.isfile", return_value=True):
                    with patch("os.access", return_value=True):
                        with patch("subprocess.run", return_value=mock_proc):
                            res = convert_html_to_markdown("<p>Fallback text</p>")
                            self.assertEqual(res, "Fallback text")

        # SimpleHtmlToMarkdownParser tags: br, a, li, empty table, table with padded cells
        html_sample = """
        <p>Строка 1<br>Строка 2</p>
        <a href="https://example.com">Ссылка</a>
        <ul>
          <li>Элемент списка</li>
        </ul>
        <table></table>
        <table><tr><td>Ячейка 1</td></tr></table>
        """
        parsed = convert_html_to_markdown(html_sample)
        self.assertIn("Строка 1", parsed)
        self.assertIn("Строка 2", parsed)
        self.assertIn("[Ссылка](https://example.com)", parsed)
        self.assertIn("- Элемент списка", parsed)
        self.assertIn("| Ячейка 1 |", parsed)

    def test_transport_edge_cases(self) -> None:
        # Empty input -> None
        t1 = StdioTransport(stdin=io.BytesIO(b""), stdout=io.BytesIO())
        self.assertIsNone(t1.read_message())

        # Whitespace line -> None
        t2 = StdioTransport(stdin=io.BytesIO(b"   \n"), stdout=io.BytesIO())
        self.assertIsNone(t2.read_message())

        # Content-Length <= 0 -> None
        header_zero = b"Content-Length: 0\r\n\r\n"
        t3 = StdioTransport(stdin=io.BytesIO(header_zero), stdout=io.BytesIO())
        self.assertIsNone(t3.read_message())

        # Multi-line header and read framed message
        valid_body = b'{"status": "ok"}'
        header_multi = b"Content-Length: " + str(len(valid_body)).encode("ascii") + b"\r\nCustom-Header: value\r\n\r\n" + valid_body
        t4 = StdioTransport(stdin=io.BytesIO(header_multi), stdout=io.BytesIO())
        msg = t4.read_message()
        self.assertIsNotNone(msg)
        self.assertEqual(msg, {"status": "ok"})
        self.assertTrue(t4.is_framed)

    def test_converter_parser_direct_and_cli_exceptions(self) -> None:
        parser = _SimpleHtmlToMarkdownParser()
        parser.feed("<p>Line 1<br>Line 2</p><a href='https://example.com'>Link</a><ul><li>item</li></ul>")
        parser.feed("<table></table>")  # empty rows
        parser.feed("<table><tr></tr></table>")  # 0 columns
        # Table with uneven rows and header padding
        parser.feed("<table><tr><th>Col1</th></tr><tr><td>A</td><td>B</td></tr></table>")
        res = parser.get_markdown()
        self.assertIn("Col1", res)
        self.assertIn("Link", res)
        self.assertIn("item", res)

        # CLI subprocess exception fallback
        with patch.object(html_markdown, "_native_h2m", None):
            with patch("shutil.which", return_value="/usr/bin/html-to-markdown"):
                with patch("os.path.isfile", return_value=True):
                    with patch("os.access", return_value=True):
                        with patch("subprocess.run", side_effect=OSError("binary crashed")):
                            res_fallback = convert_html_to_markdown("<p>Simple fallback</p>")
                            self.assertEqual(res_fallback, "Simple fallback")

    def test_rules_and_engine_edge_coverage(self) -> None:
        doc = Document()
        preset = resolve_preset("report")

        # 1. engine.py line 354: is_in_table return True in validate_paragraphs
        doc.add_paragraph("Paragraph in table test")
        with patch("gost_standardizer.core.engine.is_in_table", return_value=True):
            issues, matches = validate_paragraphs(doc, preset)
            self.assertEqual(len(issues), 0)

        # 2. rules.py line 146: r_fonts is None and added
        mock_style = MagicMock()
        mock_rpr = MagicMock()
        mock_rpr.rFonts = None
        mock_rpr._add_rFonts.return_value = MagicMock()
        mock_style._element.rPr = mock_rpr
        set_style_font(mock_style, "Arial", 12, True)
        mock_rpr._add_rFonts.assert_called_once()

        # 3. rules.py line 194: style_fmt fallback in effective_paragraph_metrics
        p = doc.add_paragraph("Testing style format fallback")
        p.paragraph_format.first_line_indent = None
        p.style.paragraph_format.first_line_indent = Pt(12)
        metrics = effective_paragraph_metrics(p)
        self.assertIsNotNone(metrics["first_line_indent_mm"])

        # 4. presets.py lines 176-177: get_profiles_dir ImportError
        with patch.dict("sys.modules", {"gost_standardizer": None}):
            p_dir = presets.get_profiles_dir()
            self.assertIsNotNone(p_dir)

    def test_transport_header_eof(self) -> None:
        # Header EOF on line 48
        t = StdioTransport(stdin=io.BytesIO(b"Custom-Header: 1\n"), stdout=io.BytesIO())
        self.assertIsNone(t.read_message())

    def test_scripts_main_execution(self) -> None:
        import runpy

        scripts_dir = Path(__file__).resolve().parents[1] / "scripts"
        mock_transport = MagicMock()
        mock_transport.read_message.return_value = None

        with patch("sys.argv", ["gost-standardizer", "presets"]):
            with patch("gost_standardizer.server.mcp_server.StdioTransport", return_value=mock_transport):
                # 1. scripts/meganorm_catalog.py as __main__
                with patch("sys.stdout", new=io.StringIO()):
                    with self.assertRaises(SystemExit) as ctx:
                        runpy.run_path(str(scripts_dir / "meganorm_catalog.py"), run_name="__main__")
                    self.assertEqual(ctx.exception.code, 0)

                # 2. scripts/mcp_server.py as __main__
                with self.assertRaises(SystemExit) as ctx:
                    runpy.run_path(str(scripts_dir / "mcp_server.py"), run_name="__main__")
                self.assertEqual(ctx.exception.code, 0)

                # 3. scripts/gost_standardizer.py as __main__
                with patch("sys.stdout", new=io.StringIO()):
                    with self.assertRaises(SystemExit) as ctx:
                        runpy.run_path(str(scripts_dir / "gost_standardizer.py"), run_name="__main__")
                    self.assertEqual(ctx.exception.code, 0)


if __name__ == "__main__":
    unittest.main()

